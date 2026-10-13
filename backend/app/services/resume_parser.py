from io import BytesIO
import re
from typing import Dict, List

from docx import Document
from pypdf import PdfReader

from app.services.resume_scorer import score_resume


class ResumeParseError(ValueError):
    """Raised when a resume file is supported but cannot be read."""


SECTION_HEADINGS = {
    "summary": {
        "summary",
        "professional summary",
        "profile",
        "career profile",
        "objective",
    },
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "work history",
    },
    "education": {"education", "academic background", "qualifications"},
    "projects": {"projects", "selected projects", "personal projects"},
    "skills": {"skills", "technical skills", "core competencies"},
}

SKILL_TERMS = [
    "Agile",
    "AWS",
    "Azure",
    "Bash",
    "C#",
    "C++",
    "CI/CD",
    "Communication",
    "CSS",
    "Data Analysis",
    "Data Visualization",
    "Django",
    "Docker",
    "Excel",
    "FastAPI",
    "Figma",
    "Flask",
    "Generative AI",
    "Git",
    "GitHub",
    "Google Cloud",
    "GraphQL",
    "HTML",
    "Java",
    "JavaScript",
    "Jest",
    "Kubernetes",
    "Leadership",
    "Linux",
    "LLM",
    "Machine Learning",
    "MongoDB",
    "MySQL",
    "NLP",
    "Node.js",
    "NumPy",
    "Pandas",
    "PHP",
    "Power BI",
    "PostgreSQL",
    "Problem Solving",
    "Product Management",
    "Project Management",
    "Pytest",
    "PyTorch",
    "Python",
    "Redis",
    "React",
    "REST API",
    "Ruby",
    "Ruby on Rails",
    "Scikit-learn",
    "Scrum",
    "SQL",
    "SQLite",
    "Spring Boot",
    "Tableau",
    "TensorFlow",
    "Terraform",
    "TypeScript",
    "Unit Testing",
    "UX Design",
    "Vue.js",
]

# Each key is the one canonical name the app stores and displays. Values list
# common ways the same skill may be written in a resume or job description.
# Keeping aliases here makes both sides of a comparison use identical names.
SKILL_ALIASES = {
    "AWS": ("Amazon Web Services",),
    "Azure": ("Microsoft Azure",),
    "C++": ("C Plus Plus",),
    "C#": ("C Sharp",),
    "Excel": ("Microsoft Excel",),
    "Google Cloud": ("Google Cloud Platform", "GCP"),
    "JavaScript": ("JS",),
    "Kubernetes": ("K8s",),
    "LLM": ("Large Language Model", "Large Language Models"),
    "Node.js": ("NodeJS", "Node JS"),
    "PostgreSQL": ("Postgres", "Postgre SQL"),
    "Power BI": ("Microsoft Power BI",),
    "React": ("React.js", "ReactJS"),
    "Ruby on Rails": ("Rails",),
    "Scikit-learn": ("Scikit learn", "sklearn"),
    "TypeScript": ("TS",),
    "Vue.js": ("Vue", "VueJS"),
}

RESUME_ENTRY_SECTIONS = ("experience", "education", "projects")
_DATE_OR_RECORD_MARKER = re.compile(r"\||\b(?:19|20)\d{2}\b|\b(?:present|current)\b", re.IGNORECASE)
_DEGREE_MARKER = re.compile(
    r"\b(?:bachelor|bsc|b\.s\.|master|msc|m\.s\.|phd|ph\.d\.|doctorate|associate|diploma|certificate)\b",
    re.IGNORECASE,
)
_ROLE_MARKER = re.compile(
    r"\b(?:engineer|developer|analyst|manager|intern|designer|architect|consultant|"
    r"scientist|assistant|specialist|coordinator|director|lead|officer)\b",
    re.IGNORECASE,
)
_BULLET_MARKER = re.compile(r"^\s*(?:[-*•●▪◦]|\d+[.)])\s+")


def _read_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ResumeParseError("This PDF is password-protected. Upload an unlocked copy.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ResumeParseError:
        raise
    except Exception as error:
        raise ResumeParseError("We couldn't read this PDF. Try exporting it again as a PDF.") from error


def _read_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    except Exception as error:
        raise ResumeParseError("We couldn't read this DOCX file. Try saving it again in Word.") from error


def detect_sections(text: str) -> Dict[str, str]:
    heading_lookup = {
        heading: section
        for section, headings in SECTION_HEADINGS.items()
        for heading in headings
    }
    found: Dict[str, List[str]] = {}
    active_section = None

    for line in text.splitlines():
        cleaned_line = re.sub(r"[^a-z0-9+# ]", "", line.lower()).strip()
        cleaned_line = re.sub(r"\s+", " ", cleaned_line)
        heading = heading_lookup.get(cleaned_line)
        if heading:
            active_section = heading
            found.setdefault(active_section, [])
            continue

        if active_section and line.strip():
            found[active_section].append(line.strip())

    return {
        section: "\n".join(lines).strip()
        for section, lines in found.items()
        if lines
    }


def _looks_like_entry_heading(section: str, line: str) -> bool:
    if _BULLET_MARKER.match(line):
        return False
    if _DATE_OR_RECORD_MARKER.search(line):
        return True
    if section == "projects" and re.search(r"\s[—–:]\s+\S", line):
        return True
    words = line.split()
    if section == "education":
        return bool(_DEGREE_MARKER.search(line))
    if section == "experience":
        if len(words) > 12 or re.search(r"[.!?]$", line):
            return False
        return bool(_ROLE_MARKER.search(line))
    return False


def _split_project_heading(line: str):
    match = re.match(r"^(.{2,100}?)\s+[—–:]\s+(.+)$", line)
    if match:
        return match.group(1).strip(), [match.group(2).strip()]
    return line, []


def extract_resume_entries(sections: Dict[str, str]) -> Dict[str, List[Dict[str, object]]]:
    """Group source lines into conservative experience, education, and project entries.

    This is formatting-based extraction: it keeps the resume's original wording and does
    not guess employer names, dates, or qualifications. Dates, separators, degree words,
    and common job-title words are used only as entry-boundary hints.
    """
    extracted: Dict[str, List[Dict[str, object]]] = {}
    for section in RESUME_ENTRY_SECTIONS:
        lines = [line.strip() for line in sections.get(section, "").splitlines() if line.strip()]
        groups: List[List[str]] = []
        current: List[str] = []
        for line in lines:
            is_heading = _looks_like_entry_heading(section, line)
            if is_heading and current:
                current_has_content = len(current) > 1
                current_has_record_marker = bool(_DATE_OR_RECORD_MARKER.search(current[0]))
                project_has_description = (
                    section == "projects" and bool(_split_project_heading(current[0])[1])
                )
                repeated_typed_heading = (
                    (section == "education" and bool(_DEGREE_MARKER.search(current[0])) and bool(_DEGREE_MARKER.search(line)))
                    or (section == "experience" and bool(_ROLE_MARKER.search(current[0])) and bool(_ROLE_MARKER.search(line)))
                )
                if (
                    current_has_content
                    or current_has_record_marker
                    or project_has_description
                    or repeated_typed_heading
                ):
                    groups.append(current)
                    current = [line]
                else:
                    # A title and its date or school line often wrap across paragraphs.
                    current.append(line)
            else:
                current.append(line)
        if current:
            groups.append(current)

        entries: List[Dict[str, object]] = []
        for group in groups:
            heading = group[0]
            details = group[1:]
            if section == "projects" and len(group) == 1:
                heading, inline_details = _split_project_heading(heading)
                details.extend(inline_details)
            entries.append({"heading": heading, "details": details})
        extracted[section] = entries
    return extracted


def _skill_pattern(name: str) -> str:
    """Build a case-insensitive token pattern, guarding dotted JS/TS suffixes."""
    # Without the dot check, the JS alias would also match the suffix in React.js.
    if name.casefold() in {"js", "ts"}:
        left_boundary = r"(?<![a-z0-9.])"
    else:
        left_boundary = r"(?<![a-z0-9])"
    return left_boundary + re.escape(name) + r"(?![a-z0-9])"


def extract_skills(text: str) -> List[str]:
    found = []
    for skill in SKILL_TERMS:
        names = (skill, *SKILL_ALIASES.get(skill, ()))
        patterns = (_skill_pattern(name) for name in names)
        if any(re.search(pattern, text, flags=re.IGNORECASE) for pattern in patterns):
            found.append(skill)
    return sorted(found, key=str.casefold)


def parse_resume(filename: str, content: bytes) -> Dict[str, object]:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        text = _read_pdf(content)
    elif extension == "docx":
        text = _read_docx(content)
    else:
        raise ResumeParseError("Please upload a PDF or DOCX resume.")

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise ResumeParseError(
            "No selectable text was found. Scanned-image PDFs need OCR, which isn't included yet."
        )

    sections = detect_sections(text)
    skills = extract_skills(text)
    score = score_resume(text, sections, skills)

    return {
        "filename": filename,
        "character_count": len(text),
        "text": text,
        "sections": sections,
        "entries": extract_resume_entries(sections),
        "skills": skills,
        **score,
    }
